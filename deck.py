"""Template engine AK3U — design token diekstrak dari 'Tugas k3 Ajie 2.pptx'."""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ASSETS = Path(__file__).parent / "assets"

NAVY = RGBColor(0x1D, 0x30, 0x5A)
GREEN = RGBColor(0x8D, 0xC6, 0x3F)
ORANGE = RGBColor(0xF5, 0x7A, 0x20)
BLUE = RGBColor(0x3E, 0x7F, 0xCD)
RED = RGBColor(0xC0, 0x39, 0x2B)
BORDER = RGBColor(0xDC, 0xE1, 0xE6)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x6E, 0x76, 0x80)
TINT = RGBColor(0xE9, 0xF1, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

# kolom konten aman: kiri 1.35" .. kanan 8.30"
X0, X1 = 1.35, 8.30
W = X1 - X0

_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
_DIAG = """<p:sp {ns}><p:nvSpPr><p:cNvPr id="{id}" name="diag{id}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
<p:spPr><a:xfrm{flip}><a:off x="{x}" y="445770"/><a:ext cx="{cx}" cy="4286250"/></a:xfrm>
<a:custGeom><a:rect b="b" l="l" r="r" t="t"/><a:pathLst><a:path extrusionOk="0" h="120000" w="120000">
<a:moveTo><a:pt x="0" y="120000"/></a:moveTo><a:lnTo><a:pt x="120000" y="0"/></a:lnTo></a:path></a:pathLst></a:custGeom>
<a:gradFill><a:gsLst><a:gs pos="0"><a:srgbClr val="3E7FCD"/></a:gs><a:gs pos="100000"><a:srgbClr val="96C0FF"/></a:gs></a:gsLst><a:lin ang="16200000" scaled="0"/></a:gradFill>
<a:ln cap="flat" cmpd="sng" w="12700"><a:solidFill><a:srgbClr val="E0E4E8"/></a:solidFill><a:prstDash val="solid"/><a:round/></a:ln>
<a:effectLst><a:outerShdw blurRad="40000" rotWithShape="0" dir="5400000" dist="23000"><a:srgbClr val="000000"><a:alpha val="34902"/></a:srgbClr></a:outerShdw></a:effectLst></p:spPr></p:sp>"""


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    return prs


def _shape(slide, shp, x, y, w, h, fill=None, line=None, adj=None):
    s = slide.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    if adj is not None:
        s.adjustments[0] = adj
    s.text_frame.word_wrap = True
    return s


def _write(tf, lines, size, color=INK, bold=False, align=PP_ALIGN.LEFT, space=2, line_spacing=0.95):
    """lines: str | list[str] | list[(text, bold, color, size)]"""
    if isinstance(lines, str):
        lines = [lines]
    tf.word_wrap = True
    for i, item in enumerate(lines):
        txt, b, c, sz = (item + (None,) * 4)[:4] if isinstance(item, tuple) else (item, None, None, None)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = txt
        f = r.font
        f.name, f.size = FONT, Pt(sz or size)
        f.bold = bold if b is None else b
        f.color.rgb = color if c is None else c
    return tf


def textbox(slide, x, y, w, h, lines, size=9, color=INK, bold=False, align=PP_ALIGN.LEFT, space=2, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.vertical_anchor = anchor
    tb.text_frame.margin_left = tb.text_frame.margin_right = Emu(45720)
    tb.text_frame.margin_top = tb.text_frame.margin_bottom = 0
    _write(tb.text_frame, lines, size, color, bold, align, space)
    return tb


def decor(slide, page=None, footer_text=None):
    """Bingkai tetap: balok navy kiri-kanan, aksen hijau, 3 diagonal gradient."""
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.075, 1.087, 1.125, 3.524, NAVY)
    _shape(slide, MSO_SHAPE.RECTANGLE, 8.735, 1.050, 1.200, 3.600, NAVY)
    _shape(slide, MSO_SHAPE.RECTANGLE, 8.399, 0.465, 1.575, 0.615, GREEN)
    _shape(slide, MSO_SHAPE.PARALLELOGRAM, 0.112, 4.349, 1.650, 0.900, GREEN, adj=0.25)
    tree = slide.shapes._spTree
    for i, (x, cx, flip) in enumerate(
        [(5623710, 685818, ' flipH="1"'), (6652437, 754400, ""), (5417965, 1577381, ' flipH="1"')]
    ):
        tree.append(etree.fromstring(_DIAG.format(ns=_NS, id=900 + i, x=x, cx=cx, flip=flip)))
    if footer_text:
        textbox(slide, 5.10, 5.29, 3.20, 0.19, footer_text, 6, MUTED, align=PP_ALIGN.RIGHT)
    if page:
        slide.shapes.add_picture(str(ASSETS / "icon.png"), Inches(8.90), Inches(0.15), Inches(0.62), Inches(0.62))


def slide_title(slide, title, sub=None):
    size = 15 if len(title) <= 48 else (13.5 if len(title) <= 56 else 12.5)
    textbox(slide, X0 - 0.02, 0.33, W, 0.30, title, size, NAVY, bold=True)
    if sub:
        textbox(slide, X0 - 0.02, 0.63, W, 0.22, sub, 8, MUTED)


def pill(slide, x, y, w, text, color=NAVY, size=7.5, h=0.24):
    s = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, color, adj=0.25)
    s.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    s.text_frame.margin_left = s.text_frame.margin_right = Emu(45720)
    s.text_frame.margin_top = s.text_frame.margin_bottom = 0
    _write(s.text_frame, text, size, WHITE, True, PP_ALIGN.CENTER, 0)
    return s


def card(slide, x, y, w, h, header=None, body=None, accent=NAVY, size=8, head_size=7.5, align=PP_ALIGN.LEFT):
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, BORDER, adj=0.06)
    ty = y + 0.09
    if header:
        pill(slide, x + 0.09, ty, w - 0.18, header, accent, head_size)
        ty += 0.32
    if body:
        textbox(slide, x + 0.12, ty, w - 0.24, y + h - ty - 0.08, body, size, align=align)


def note(slide, x, y, w, text, color=TINT, fg=NAVY, size=7.5, h=0.26):
    s = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, color, adj=0.25)
    s.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    s.text_frame.margin_top = s.text_frame.margin_bottom = 0
    _write(s.text_frame, text, size, fg, True, PP_ALIGN.CENTER, 0)


def numbered(slide, x, y, w, items, row_h=0.30, gap=0.04, size=8, accent=NAVY, start=1):
    for i, it in enumerate(items):
        ry = y + i * (row_h + gap)
        _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, ry, w, row_h, RGBColor(0xF4, 0xF7, 0xFA), BORDER, adj=0.2)
        pill(slide, x + 0.09, ry + (row_h - 0.21) / 2, 0.30, str(start + i), accent, size - 0.5, h=0.21)
        textbox(slide, x + 0.50, ry, w - 0.60, row_h, it, size, anchor=MSO_ANCHOR.MIDDLE)


def table(slide, x, y, w, h, rows, widths=None, size=7, head_size=7, accent=NAVY, row_h=0.24):
    nrow, ncol = len(rows), len(rows[0])
    g = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w), Inches(h))
    t = g.table
    t.first_row = True
    if widths:
        for i, fr in enumerate(widths):
            t.columns[i].width = Emu(int(Inches(w) * fr / sum(widths)))
    for ri, row in enumerate(rows):
        t.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.margin_left = cell.margin_right = Emu(54000)
            cell.margin_top = cell.margin_bottom = Emu(9000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = accent if ri == 0 else (WHITE if ri % 2 else RGBColor(0xF4, 0xF7, 0xFA))
            txt, color = (val if isinstance(val, tuple) else (val, None))
            _write(
                cell.text_frame,
                str(txt),
                head_size if ri == 0 else size,
                WHITE if ri == 0 else (color or INK),
                ri == 0,
                PP_ALIGN.CENTER if (ri == 0 or ci == 0) else PP_ALIGN.LEFT,
                0,
            )
    return t


def picture(slide, path, x, y, w, h, caption=None):
    """Gambar di-fit (contain) ke dalam kotak, rasio asli dipertahankan."""
    from PIL import Image

    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h + (0.20 if caption else 0), WHITE, BORDER, adj=0.05)
    iw, ih = Image.open(ASSETS / path).size
    bw, bh = w - 0.12, h - 0.12
    sc = min(bw / iw, bh / ih)
    pw, ph = iw * sc, ih * sc
    slide.shapes.add_picture(str(ASSETS / path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    if caption:
        textbox(slide, x, y + h, w, 0.18, caption, 6, MUTED, align=PP_ALIGN.CENTER)


def cover(prs, kelompok, judul, unit_count, durasi, asesi="Mohammad Ajie Wicaksono"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    decor(s, page=True)
    _shape(s, MSO_SHAPE.OVAL, 0.41, 1.09, 1.80, 1.80, WHITE, BORDER)
    s.shapes.add_picture(str(ASSETS / "keyvisual.png"), Inches(0.55), Inches(1.17), Inches(1.54), Inches(1.54))
    textbox(s, 2.50, 0.80, 5.30, 0.70, ["AHLI KESELAMATAN DAN KESEHATAN", "KERJA UMUM (AK3U)"], 17, NAVY, bold=True, space=0)
    pill(s, 2.52, 1.60, 2.90, f"TUGAS KELOMPOK PEKERJAAN {kelompok}", NAVY, 9, h=0.28)
    textbox(s, 2.52, 2.02, 5.30, 0.60, judul, 13, NAVY, bold=True)
    textbox(s, 2.52, 2.70, 5.30, 0.40, "FR.IA.04A — Daftar Instruksi Terstruktur (Penjelasan Proyek Singkat)", 8.5, MUTED)
    _shape(s, MSO_SHAPE.RECTANGLE, 2.55, 3.10, 1.60, 0.03, GREEN)
    textbox(
        s, 2.52, 3.28, 5.30, 0.70,
        [("Skema Sertifikasi : AHLI KESELAMATAN DAN KESEHATAN KERJA UMUM", False, MUTED, 8),
         ("Nomor Skema : SKM/2037/00013/2/2021/1", False, MUTED, 8),
         (f"Lingkup : {unit_count} unit kompetensi  |  Presentasi {durasi}", False, MUTED, 8)], 8,
    )
    textbox(s, 2.52, 4.08, 5.30, 0.24, f"Oleh: {asesi}", 10, NAVY, bold=True)
    return s


def content(prs, title, sub=None, kelompok=None, page=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    foot = f"Daftar Instruksi Terstruktur AK3U BNSP  |  Kelompok Pekerjaan {kelompok}  |  {page}" if kelompok else None
    decor(s, footer_text=foot)
    slide_title(s, title, sub)
    return s
